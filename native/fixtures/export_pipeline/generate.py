"""生成导出 pipeline 逐字节对照夹具（rust-native-core 任务 4.3）——**可选的历史 golden 生成器**。
生成的 12 个 golden 已固化在 `native/fixtures/export_pipeline/golden/`，
Rust 侧 `tests/compat/tests/export_pipeline.rs` 逐字节比对的是这些文件，不需要 Python；
只有改动产物格式后要重造参照 golden 时才需要仓库内 Python。

运行方式（仓库根目录）::

    ct/.venv/Scripts/python.exe native/fixtures/export_pipeline/generate.py

产物：
- workspace/ —— 完整输入工作区（schema YAML、Excel 数据工作簿、manifest、i18n）
- golden/ —— Python run_pipeline 的全部产物字节（output/** + layout manifest）
Rust 测试把 workspace/ 复制到临时目录跑 Rust pipeline，逐文件与 golden 对照。
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent.parent.parent / "ct" / "src"))

from ct.app.exporting.build import run_pipeline  # noqa: E402
from ct.app.exporting.models import ExportRequest  # noqa: E402
from ct.excel.layout import LayoutBuilder  # noqa: E402
from ct.excel.canonical_template import build_canonical_template  # noqa: E402
from ct.excel.layout_manifest import LayoutManifest, manifest_payload  # noqa: E402
from ct.schema.hashing import compute_schema_hash  # noqa: E402

WORKSPACE = HERE / "workspace"
GOLDEN = HERE / "golden"

GLOBAL_YAML = "primary_lang: zh\nsecondary_langs:\n  - en\n"

ITEM_YAML = """table: Item
primary: Id
fields:
  - name: Id
    type: int32
  - name: CodeName
    type: string
  - name: Name
    type: string
    i18n: true
  - name: Quality
    type: Rarity
  - name: Price
    type: int32
indexes:
  - kind: codename
"""

RARITY_YAML = """kind: enum
name: Rarity
values:
  - name: Common
  - name: Rare
    comment: 稀有
"""

ROWS = [
    {"Id": 1001, "CodeName": "sword", "Name": "铁剑", "Quality": "Common", "Price": 320},
    {"Id": 1002, "CodeName": "shield", "Name": "铁盾", "Quality": "Rare", "Price": 480},
]

I18N_EN = {
    "1001.Name": {"text": "Iron Sword", "confirmed": True, "status": "confirmed", "source": "铁剑"},
    "1002.Name": {"text": "", "confirmed": False, "status": "missing", "source": "铁盾"},
}


def write(path: Path, content) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_bytes(content.encode("utf-8"))


def main() -> None:
    if WORKSPACE.exists():
        shutil.rmtree(WORKSPACE)
    if GOLDEN.exists():
        shutil.rmtree(GOLDEN)
    write(WORKSPACE / "config/global.yaml", GLOBAL_YAML)
    write(WORKSPACE / "config/schemas/item.yaml", ITEM_YAML)
    write(WORKSPACE / "config/types/rarity.yaml", RARITY_YAML)

    # 模板 + manifest（schema_hash 用真实 canonical hash，保证读取兼容性闸门通过）
    from ct.config import load_config
    from ct.schema.resource_repository import YamlResourceRepository

    config = load_config(WORKSPACE)
    repo = YamlResourceRepository(config.resolve("schemas_dir"), config.resolve("types_dir"))
    ws = repo.load()
    table = ws.tables[0]
    enums = {e.name: e for e in ws.enums}
    schema_hash = compute_schema_hash(table, tuple(ws.resources))
    layout = LayoutBuilder(table, schema_hash=schema_hash, records={}).build()
    write(
        WORKSPACE / "excel/Item.xlsx",
        build_canonical_template(layout, enums=enums, primary=table.primary),
    )
    manifest = LayoutManifest.from_layout(layout)
    write(WORKSPACE / "excel/layout_manifests/Item.json", manifest_payload(manifest))

    # 数据行（openpyxl 写入表头之后）
    from openpyxl import load_workbook

    wb = load_workbook(str(WORKSPACE / "excel/Item.xlsx"))
    ws1 = wb.active
    header = [ws1.cell(row=2, column=i + 1).value for i in range(len(ROWS[0]))]
    # 表头第二行是 "Name（type）" 形态；按列序直接写
    for r, row in enumerate(ROWS, start=3):
        for c, value in enumerate(row.values(), start=1):
            ws1.cell(row=r, column=c, value=value)
    wb.save(str(WORKSPACE / "excel/Item.xlsx"))

    # i18n 文件（主语言源来自 Excel，不单独成文件）
    write(
        WORKSPACE / "i18n/en/Item.json",
        json.dumps(I18N_EN, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
    )

    # Python 全量导出 → golden
    result = run_pipeline(ExportRequest(root=WORKSPACE))
    for path in sorted(WORKSPACE.rglob("*")):
        if not path.is_file():
            continue
        rel = path.relative_to(WORKSPACE)
        if rel.parts[0] in {"output", "excel"} and "Item.xlsx" not in rel.name:
            write(GOLDEN / rel, path.read_bytes())
    print("tables:", result.tables, "langs:", result.languages)
    print("written:", len(result.written), "reused:", len(result.reused))
    print("bundle_hashes:", {k: v[:12] for k, v in result.bundle_hashes.items()})
    print("wrote golden to", GOLDEN)


if __name__ == "__main__":
    main()
