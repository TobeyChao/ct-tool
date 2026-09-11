"""N3 配套：把旧版 Excel 模板升级到当前 ct-tool 的模板格式（保留数据）。

背景：
  游戏里的 `Config/gd/excel/*.xlsx` 是**旧版生成器**产出的，表头顺序是
    r1 = `Name\\ntype`  →  r2 = record 子类型行（仅含 record 的表）  →  r3 = 注释行  →  数据
  而当前 ct-tool 的模板是
    r1 = 注释  →  r2 = `Name\\ntype`  →  r3 = 子注释  →  r4 = 子类型  →  数据
  两者 `header_rows` 不同（Item 是 3 vs 4）⇒ 直接导出会把**第一行数据当表头吃掉**
  （实测：Item 只剩 Id=2,3,4，并连带触发 Quest 的外键校验失败）。

做法：重新生成当前格式的模板，再把旧 Excel 的**数据行**按列位置搬过去；
      顺带把 `vector<T>` 单元格从旧的 `a,b,c` 文法改成 `[a,b,c]`。

用法：
  CT_TOOL=/path/to/ct-tool ct/.venv/bin/python migrate_excel.py <workspace_dir> [--src <旧的 excel 目录>]
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
CT = HERE.parents[1]
sys.path.insert(0, str(CT / "ct" / "src"))

from openpyxl import load_workbook  # noqa: E402

from ct.config import load_config  # noqa: E402
from ct.excel.canonical_template import generate_canonical_template  # noqa: E402
from ct.excel.layout import build_layout  # noqa: E402
from ct.schema.hashing import compute_schema_hash  # noqa: E402
from ct.schema.resource_repository import YamlResourceRepository  # noqa: E402
from ct.schema.type_expression import VectorType  # noqa: E402


def _is_int_cell(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _data_start(sh) -> int:
    """旧格式的数据起始行：第一行「第 1 列是数字」的行。"""
    for i, row in enumerate(sh.iter_rows(values_only=True), 1):
        if row and _is_int_cell(row[0]):
            return i
    return sh.max_row + 1


def _fix_vector_cell(value, type_text: str):
    """旧文法 `a,b,c` → 新文法 `[a,b,c]`（已经是 [...] 或空值则原样返回）。"""
    if value is None:
        return value
    text = str(value).strip()
    if text == "" or (text.startswith("[") and text.endswith("]")):
        return value
    parts = [p.strip() for p in re.split(r"[,\uFF0C]", text) if p.strip()]
    joined = ",".join(parts)
    if type_text.startswith("vector<string>"):
        joined = ",".join(f'"{p}"' for p in parts)
    return f"[{joined}]"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workspace", type=Path)
    ap.add_argument("--src", type=Path, default=None, help="旧 excel 目录（默认取 workspace/excel）")
    ap.add_argument("--in-place", action="store_true", help="核对无误后直接覆盖 workspace/excel")
    args = ap.parse_args()

    ws = args.workspace.resolve()
    src_dir = (args.src or (ws / "excel")).resolve()
    cfg = load_config(ws)
    res = YamlResourceRepository(cfg.resolve("schemas_dir"), cfg.resolve("types_dir")).load()
    records = {r.name: r for r in res.records}
    enums = {e.name: e for e in res.enums}

    tmp = Path(tempfile.mkdtemp(prefix="excel-migrate-"))
    for table in res.tables:
        layout = build_layout(
            table,
            schema_hash=compute_schema_hash(table, tuple(records.values())),
            records=records,
        )
        old_path = src_dir / f"{table.table}.xlsx"
        new_path = tmp / f"{table.table}.xlsx"
        generate_canonical_template(layout, new_path, enums=enums, primary=table.primary)

        old_sh = load_workbook(old_path, data_only=True).active
        start = _data_start(old_sh)
        new_wb = load_workbook(new_path)
        new_sh = new_wb.active
        target = layout.header_rows + 1

        vector_cols = {
            c.index: c.type_text
            for c in layout.columns
            if isinstance(_field_type(table, c), VectorType)
        }
        moved = 0
        for offset, row in enumerate(
            old_sh.iter_rows(min_row=start, values_only=True)
        ):
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
                continue
            r = target + moved
            for col_index, value in enumerate(row, 1):
                if value is None:
                    continue
                if col_index in vector_cols:
                    value = _fix_vector_cell(value, vector_cols[col_index])
                new_sh.cell(row=r, column=col_index, value=value)
            moved += 1
        new_wb.save(new_path)
        print(f"  {table.table:12} 旧数据起于 r{start} → 新数据起于 r{target}，搬了 {moved} 行")

    if args.in_place:
        dest = ws / "excel"
        for f in tmp.glob("*.xlsx"):
            shutil.copy2(f, dest / f.name)
        print(f"\n✅ 已覆盖 {dest}")
    else:
        print(f"\n迁移后的 Excel 在: {tmp}")
        print("（默认不覆盖 workspace/excel；核对后加 --in-place 直接应用）")


def _field_type(table, column):
    """按 stable_path 找到该列的 TypeExpression（只处理顶层字段）。"""
    top = column.stable_path[len(table.resource_id) + 1:].split("/")[0]
    top = top.split("[")[0]
    for field in table.fields:
        if field.name == top:
            return field.type_expr
    return None


if __name__ == "__main__":
    main()
