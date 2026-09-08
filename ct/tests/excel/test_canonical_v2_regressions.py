from pathlib import Path

from openpyxl import Workbook, load_workbook

from ct.excel.canonical_reader import parse_vector_cell, read_canonical_excel
from ct.excel.canonical_template import generate_canonical_template
from ct.excel.layout import build_layout
from ct.schema.resources import FieldDef, RecordResource, TableResource


def test_bracket_vector_parser_is_typed_and_strict() -> None:
    assert parse_vector_cell('[1, 2, 3]', 'int32') == ([1, 2, 3], None)
    assert parse_vector_cell('["a,b", "x\\ny"]', 'string') == (["a,b", "x\ny"], None)
    assert parse_vector_cell('1,2', 'int32')[1]
    assert parse_vector_cell('[1,]', 'int32')[1]


def test_fixed_record_vector_uses_last_filled_slot_and_defaults(tmp_path: Path) -> None:
    record = RecordResource(name='Reward', fields=[FieldDef(name='Min', type='int32'), FieldDef(name='Max', type='int32')])
    table = TableResource(table='Item', primary='Id', fields=[FieldDef(name='Id', type='int32'), FieldDef(name='Rewards', type='vector<Reward>', excel_columns=3)])
    layout = build_layout(table, schema_hash='x', records={'Reward': record})
    path = tmp_path / 'Item.xlsx'
    wb = Workbook(); ws = wb.active
    for _ in range(layout.header_rows): ws.append([None] * layout.column_count)
    ws.append([1, None, None, None, None, 30, None])
    wb.save(path)
    parsed = read_canonical_excel(path, layout, table, records={'Reward': record})
    assert parsed.rows == [{'Id': 1, 'Rewards': [{'Min': 0, 'Max': 0}, {'Min': 0, 'Max': 0}, {'Min': 30, 'Max': 0}]}]


def test_template_uses_full_data_region_and_no_filter(tmp_path: Path) -> None:
    table = TableResource(table='Item', primary='Id', fields=[FieldDef(name='Id', type='int32'), FieldDef(name='Tags', type='vector<int32>')])
    layout = build_layout(table, schema_hash='x', records={})
    path = tmp_path / 'Item.xlsx'
    generate_canonical_template(layout, path, enums={}, primary='Id')
    ws = load_workbook(path).active
    assert ws.freeze_panes == 'A3'
    assert ws.auto_filter.ref is None
    assert any('1048576' in str(dv.sqref) for dv in ws.data_validations.dataValidation)
