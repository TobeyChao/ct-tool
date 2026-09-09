"""Excel data change-planning tests (4.5-4.7)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText

from ct.app.canonical_commands import (
    _migrate_excel_rows,
    _publish_staged_workbook,
    _validate_staged_workbook,
)
from ct.excel.canonical_template import generate_canonical_template
from ct.excel.layout import build_layout
from ct.excel.layout_manifest import LayoutManifest
from ct.excel.planning import plan_excel_migration


def _tracked_manifest(layout):
    return LayoutManifest.from_layout(layout)
from ct.schema.resources import FieldDef, RecordResource, TableResource


def _make_workbook(path: Path, header_rows: int, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    for _ in range(header_rows):
        ws.append([""] * 8)
    for row in rows:
        ws.append(row)
    wb.save(str(path))


def _layout(table: TableResource, records: dict[str, RecordResource]):
    return build_layout(table, schema_hash="s", records=records)


def _records() -> dict[str, RecordResource]:
    return {
        "DropReward": RecordResource(
            name="DropReward",
            fields=[FieldDef(name="ItemId", type="int32"), FieldDef(name="Count", type="int32")],
        ),
    }


def test_migrate_rows_preserves_values_when_columns_move(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Name", type="string")],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Name", type="string"), FieldDef(name="Id", type="int32"), FieldDef(name="Price", type="int32")],
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    old_path = tmp_path / "old.xlsx"
    new_path = tmp_path / "new.xlsx"
    generate_canonical_template(old_layout, old_path, enums={}, primary="Id")
    workbook = load_workbook(old_path)
    worksheet = workbook.active
    worksheet.cell(old_layout.header_rows + 1, 1).value = 7
    worksheet.cell(old_layout.header_rows + 1, 2).value = "剑"
    workbook.save(old_path)
    workbook.close()
    generate_canonical_template(new_layout, new_path, enums={}, primary="Id")

    _migrate_excel_rows(
        old_path, new_path, old_layout, new_layout, _tracked_manifest(old_layout)
    )

    workbook = load_workbook(new_path, data_only=True)
    values = tuple(next(workbook.active.iter_rows(
        min_row=new_layout.header_rows + 1,
        max_row=new_layout.header_rows + 1,
        values_only=True,
    )))
    workbook.close()
    assert values == ("剑", 7, None)


def test_migrate_rows_preserves_rich_text_headers(tmp_path: Path) -> None:
    table = TableResource(
        table="Item",
        primary="Id",
        fields=[FieldDef(name="Id", type="int32")],
    )
    layout = _layout(table, {})
    old_path = tmp_path / "old.xlsx"
    new_path = tmp_path / "new.xlsx"
    generate_canonical_template(layout, old_path, enums={}, primary="Id")
    generate_canonical_template(layout, new_path, enums={}, primary="Id")

    workbook = load_workbook(old_path)
    workbook.active.cell(layout.header_rows + 1, 1).value = 7
    workbook.save(old_path)
    workbook.close()

    _migrate_excel_rows(
        old_path, new_path, layout, layout, _tracked_manifest(layout)
    )

    workbook = load_workbook(new_path, rich_text=True)
    assert isinstance(workbook.active.cell(2, 1).value, CellRichText)
    assert workbook.active.cell(layout.header_rows + 1, 1).value == 7
    workbook.close()


def test_validate_and_publish_staged_workbook(tmp_path: Path) -> None:
    target = tmp_path / "item.xlsx"
    staged = tmp_path / ".item.staged.xlsx"
    _make_workbook(target, 1, [["old"]])
    _make_workbook(staged, 1, [["new"]])

    _validate_staged_workbook(staged)
    _publish_staged_workbook(staged, target)

    assert not staged.exists()
    workbook = load_workbook(target, read_only=True)
    assert workbook.active.cell(2, 1).value == "new"
    workbook.close()


def test_validate_staged_workbook_rejects_invalid_zip(tmp_path: Path) -> None:
    staged = tmp_path / ".item.staged.xlsx"
    staged.write_bytes(b"not an xlsx")

    try:
        _validate_staged_workbook(staged)
    except ValueError as exc:
        assert "候选文件校验失败" in str(exc)
    else:
        raise AssertionError("invalid XLSX should be rejected")


def test_migrate_rows_uses_new_header_row_count(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id", fields=[FieldDef(name="Id", type="int32")]
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Drop", type="DropRange")],
    )
    records = {
        "DropRange": RecordResource(
            name="DropRange",
            fields=[FieldDef(name="Min", type="int32")],
        )
    }
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, records)
    assert old_layout.header_rows == 2
    assert new_layout.header_rows == 4
    old_path = tmp_path / "old.xlsx"
    new_path = tmp_path / "new.xlsx"
    generate_canonical_template(old_layout, old_path, enums={}, primary="Id")
    workbook = load_workbook(old_path)
    workbook.active.cell(old_layout.header_rows + 1, 1).value = 7
    workbook.save(old_path)
    workbook.close()
    generate_canonical_template(new_layout, new_path, enums={}, primary="Id")

    _migrate_excel_rows(
        old_path, new_path, old_layout, new_layout, _tracked_manifest(old_layout)
    )

    workbook = load_workbook(new_path, data_only=True)
    worksheet = workbook.active
    assert worksheet.cell(new_layout.header_rows + 1, 1).value == 7
    assert worksheet.cell(new_layout.header_rows, 1).value != 7
    workbook.close()


def test_reorder_maps_by_stable_path_not_position(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[
            FieldDef(name="Id", type="int32"),
            FieldDef(name="Name", type="string"),
            FieldDef(name="Price", type="int32"),
        ],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[
            FieldDef(name="Id", type="int32"),
            FieldDef(name="Price", type="int32"),
            FieldDef(name="Name", type="string"),
        ],
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, "剑", 10]])

    plan = plan_excel_migration(
        old_layout, new_layout, tmp_path / "item.xlsx",
        manifest=_tracked_manifest(old_layout),
    )
    assert plan.blocked is False
    by_old = {migration.old_path: migration for migration in plan.migrations}
    assert by_old["table:Item/Name"].new_index == 3  # old col2 -> new col3
    assert by_old["table:Item/Price"].new_index == 2  # old col3 -> new col2
    assert by_old["table:Item/Id"].new_index == 1


def test_explicit_rename_map_wins(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Name", type="string")],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="DisplayName", type="string")],
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, "剑"]])

    plan = plan_excel_migration(
        old_layout,
        new_layout,
        tmp_path / "item.xlsx",
        rename_map={"table:Item/Name": "table:Item/DisplayName"},
        manifest=_tracked_manifest(old_layout),
    )
    assert plan.blocked is False
    name = next(m for m in plan.migrations if m.old_path == "table:Item/Name")
    assert name.new_path == "table:Item/DisplayName"


def test_deleted_column_with_data_blocks(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Legacy", type="string")],
    )
    new_table = TableResource(
        table="Item", primary="Id", fields=[FieldDef(name="Id", type="int32")]
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, "老数据"]])

    plan = plan_excel_migration(
        old_layout, new_layout, tmp_path / "item.xlsx",
        manifest=_tracked_manifest(old_layout),
    )
    assert plan.blocked is True
    blocker = next(issue for issue in plan.issues if issue.kind == "blocker")
    assert "Legacy" in blocker.message
    assert blocker.samples == ("老数据",)


def test_type_conversion_failure_blocks(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Level", type="string")],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Level", type="int32")],
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, "abc"]])

    plan = plan_excel_migration(
        old_layout, new_layout, tmp_path / "item.xlsx",
        manifest=_tracked_manifest(old_layout),
    )
    blocker = next(issue for issue in plan.issues if issue.kind == "blocker")
    assert "不可转换" in blocker.message
    assert blocker.samples == ("abc",)


def test_enum_removal_with_data_blocks(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Rarity", type="ItemRarity")],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[FieldDef(name="Id", type="int32"), FieldDef(name="Rarity", type="ItemRarity")],
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, "Legendary"]])

    plan = plan_excel_migration(
        old_layout,
        new_layout,
        tmp_path / "item.xlsx",
        old_enums={"ItemRarity": ("Common", "Rare", "Legendary")},
        new_enums={"ItemRarity": ("Common", "Rare")},
        manifest=_tracked_manifest(old_layout),
    )
    blocker = next(issue for issue in plan.issues if issue.kind == "blocker")
    assert "Legendary" in blocker.message


def test_excel_columns_shrink_with_data_blocks(tmp_path: Path) -> None:
    records = _records()
    old_table = TableResource(
        table="Item", primary="Id",
        fields=[
            FieldDef(name="Id", type="int32"),
            FieldDef(name="Rewards", type="vector<DropReward>", excel_columns=3),
        ],
    )
    new_table = TableResource(
        table="Item", primary="Id",
        fields=[
            FieldDef(name="Id", type="int32"),
            FieldDef(name="Rewards", type="vector<DropReward>", excel_columns=1),
        ],
    )
    old_layout = _layout(old_table, records)
    new_layout = _layout(new_table, records)
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1, None, None, 30, 1, None, None]])

    plan = plan_excel_migration(
        old_layout, new_layout, tmp_path / "item.xlsx",
        manifest=_tracked_manifest(old_layout),
    )
    blocker = next(issue for issue in plan.issues if issue.kind == "blocker")
    assert "Rewards" in blocker.message


def test_untracked_workbook_requires_review(tmp_path: Path) -> None:
    old_table = TableResource(
        table="Item", primary="Id", fields=[FieldDef(name="Id", type="int32")]
    )
    new_table = TableResource(
        table="Item", primary="Id", fields=[FieldDef(name="Id", type="int32")]
    )
    old_layout = _layout(old_table, {})
    new_layout = _layout(new_table, {})
    _make_workbook(tmp_path / "item.xlsx", old_layout.header_rows, [[1]])

    plan = plan_excel_migration(old_layout, new_layout, tmp_path / "item.xlsx", manifest=None)
    assert plan.untracked is True
    assert plan.blocked is True
    assert any(issue.kind == "untracked" for issue in plan.issues)
