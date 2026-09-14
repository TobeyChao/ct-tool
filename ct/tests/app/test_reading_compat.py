"""Reading-compatibility gate tests (tasks 4.1-4.3)."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ct.app.canonical_commands import canonical_gen_template, canonical_validate
from ct.app.canonical_export import run_canonical_export
from ct.app.canonical_workspace import CanonicalWorkspace
from ct.excel.layout import build_layout
from ct.excel.layout_manifest import load_manifest
from ct.excel.reading_compat import check_reading_compatibility, expected_header_grid
from ct.schema.hashing import compute_schema_hash
from _helpers import build_project, write_yaml


def _item_schema(fields=None) -> dict:
    return {
        "table": "Item",
        "primary": "Id",
        "fields": fields
        or [
            {"name": "Id", "type": "int32"},
            {"name": "Name", "type": "string", "comment": "名称"},
            {"name": "Level", "type": "int32", "comment": "等级"},
        ],
    }


def _workspace(tmp_path: Path, *, types: list[dict] | None = None) -> Path:
    return build_project(
        tmp_path / "gd",
        schemas=[_item_schema()],
        types=types or [],
    )


def _gate(root: Path, table: str = "Item"):
    workspace = CanonicalWorkspace.load(root)
    resource = next(item for item in workspace.tables if item.table == table)
    layout = build_layout(
        resource,
        schema_hash=compute_schema_hash(resource, (*workspace.records, *workspace.enums)),
        records={record.name: record for record in workspace.records},
    )
    manifest = load_manifest(workspace.resolve("excel_dir") / "layout_manifests", table)
    excel_path = workspace.resolve("excel_dir") / (resource.excel_file or f"{table}.xlsx")
    return check_reading_compatibility(table, layout, excel_path, manifest=manifest)


def _write_rows(root: Path, table: str, rows: list[tuple]) -> None:
    path = root / "excel" / f"{table}.xlsx"
    workbook = load_workbook(path)
    sheet = workbook.active
    for row_offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=3 + row_offset, column=column, value=value)
    workbook.save(path)


def test_generated_template_is_readable(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    result = _gate(root)
    assert result.ok, result.issues


def test_expected_grid_matches_the_written_template(tmp_path: Path) -> None:
    """The gate's expectation must not drift from what the template writes."""
    root = _workspace(
        tmp_path,
        types=[
            {
                "kind": "record",
                "name": "DropReward",
                "fields": [
                    {"name": "Min", "type": "int32"},
                    {"name": "Max", "type": "int32"},
                ],
            }
        ],
    )
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Rewards", "type": "vector<DropReward>", "excel_columns": 2},
            ]
        ),
    )
    canonical_gen_template(root, all_tables=True)
    workspace = CanonicalWorkspace.load(root)
    table = next(item for item in workspace.tables if item.table == "Item")
    layout = build_layout(
        table,
        schema_hash=compute_schema_hash(table, (*workspace.records, *workspace.enums)),
        records={record.name: record for record in workspace.records},
    )
    grid = expected_header_grid(layout)
    workbook = load_workbook(root / "excel" / "Item.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    actual = {
        (row, column): str(value)
        for row, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=layout.header_rows, values_only=True), start=1
        )
        for column, value in enumerate(values, start=1)
        if value is not None and str(value).strip()
    }
    workbook.close()
    assert all(actual.get(key) == text for key, text in grid.items())
    assert not result_blocked(root)


def result_blocked(root: Path) -> bool:
    return not _gate(root).ok


def test_same_type_column_reorder_blocks(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    assert _gate(root).ok

    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Level", "type": "int32", "comment": "等级"},
                {"name": "Name", "type": "string", "comment": "名称"},
            ]
        ),
    )
    result = _gate(root)
    assert not result.ok
    assert result.reason in {"manifest-layout", "workbook-headers"}


def test_field_rename_blocks(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "DisplayName", "type": "string", "comment": "名称"},
                {"name": "Level", "type": "int32", "comment": "等级"},
            ]
        ),
    )
    assert not _gate(root).ok


def test_vector_slot_change_blocks(tmp_path: Path) -> None:
    root = _workspace(
        tmp_path,
        types=[
            {
                "kind": "record",
                "name": "DropReward",
                "fields": [{"name": "Min", "type": "int32"}],
            }
        ],
    )
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Rewards", "type": "vector<DropReward>", "excel_columns": 2},
            ]
        ),
    )
    canonical_gen_template(root, all_tables=True)
    assert _gate(root).ok

    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Rewards", "type": "vector<DropReward>", "excel_columns": 3},
            ]
        ),
    )
    result = _gate(root)
    assert not result.ok
    assert result.reason == "manifest-layout"


def test_missing_manifest_blocks(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    shutil.rmtree(root / "excel" / "layout_manifests")
    result = _gate(root)
    assert not result.ok
    assert result.reason == "manifest-missing"


def test_manifest_cannot_prove_workbook_structure(tmp_path: Path) -> None:
    """Manifest regenerated, but the workbook itself is the old one."""
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    stale = tmp_path / "stale-Item.xlsx"
    shutil.copy2(root / "excel" / "Item.xlsx", stale)

    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Name", "type": "string", "comment": "名称"},
                {"name": "Level", "type": "int32", "comment": "等级"},
                {"name": "Score", "type": "float", "comment": "分数"},
            ]
        ),
    )
    # 只重建 manifest（伪造 gen-template 的 manifest 阶段），工作簿保持旧的
    from ct.excel.layout_manifest import save_manifest

    workspace = CanonicalWorkspace.load(root)
    table = next(item for item in workspace.tables if item.table == "Item")
    layout = build_layout(
        table,
        schema_hash=compute_schema_hash(table, (*workspace.records, *workspace.enums)),
        records={record.name: record for record in workspace.records},
    )
    save_manifest(
        workspace.resolve("excel_dir") / "layout_manifests",
        "Item",
        type(load_manifest(workspace.resolve("excel_dir") / "layout_manifests", "Item")).from_layout(layout),
    )
    result = _gate(root)
    assert not result.ok
    assert result.reason == "workbook-headers"


def test_cosmetic_comment_change_still_reads(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "Name", "type": "string", "comment": "改过的注释"},
                {"name": "Level", "type": "int32", "comment": "等级"},
            ]
        ),
    )
    assert _gate(root).ok


# --------------------------------------------------------------------------- #
# 4.2 validate / export wiring
# --------------------------------------------------------------------------- #


def _with_ref_workspace(tmp_path: Path) -> Path:
    root = build_project(
        tmp_path / "refgd",
        schemas=[
            {"table": "ItemType", "primary": "Id", "fields": [{"name": "Id", "type": "int32"}]},
            {
                "table": "Item",
                "primary": "Id",
                "fields": [
                    {"name": "Id", "type": "int32"},
                    {"name": "TypeId", "type": "int32", "ref": "ItemType.Id"},
                ],
            },
        ],
    )
    return root


def test_validate_reports_template_blocker_before_reading(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "DisplayName", "type": "string", "comment": "名称"},
                {"name": "Level", "type": "int32", "comment": "等级"},
            ]
        ),
    )
    issues = canonical_validate(root)
    assert issues, "布局不兼容时必须阻止读取"
    assert any("manifest" in issue.message or "表头" in issue.message for issue in issues)


def test_export_failure_keeps_artifacts_and_manifest(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    _write_rows(root, "Item", [(1, "x", 2)])
    manifest_dir = root / "excel" / "layout_manifests"
    run_canonical_export(root)
    manifest_before = {path.name: path.read_bytes() for path in manifest_dir.glob("*.json")}
    output_before = {
        str(path.relative_to(root)): path.read_bytes()
        for path in (root / "output").rglob("*")
        if path.is_file()
    }
    ledger = root / "cache" / "state.json"
    ledger_before = ledger.read_bytes() if ledger.exists() else None

    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "DisplayName", "type": "string", "comment": "名称"},
                {"name": "Level", "type": "int32", "comment": "等级"},
            ]
        ),
    )
    from ct.app.exporting.models import ExportRequest
    from ct.app.exporting.service import run_export

    with pytest.raises(Exception):
        run_export(ExportRequest(root=root, forced=True))

    assert {path.name: path.read_bytes() for path in manifest_dir.glob("*.json")} == manifest_before
    assert {
        str(path.relative_to(root)): path.read_bytes()
        for path in (root / "output").rglob("*")
        if path.is_file()
    } == output_before
    assert (ledger.read_bytes() if ledger.exists() else None) == ledger_before


def test_filtered_export_reads_reference_dependency(tmp_path: Path) -> None:
    root = _with_ref_workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    _write_rows(root, "ItemType", [(1,)])
    _write_rows(root, "Item", [(1, 1)])

    from ct.app.exporting.models import ExportRequest
    from ct.app.exporting.service import run_export

    result = run_export(ExportRequest(root=root, table_filter="Item", forced=True))
    assert result is not None
    # 依赖表被读取用于外键校验，但不产生产物
    assert not (root / "output" / "json" / "ItemType_zh.json").exists()
    assert (root / "output" / "json" / "Item_zh.json").exists()

    _write_rows(root, "Item", [(1, 99)])
    with pytest.raises(Exception):
        run_export(ExportRequest(root=root, table_filter="Item", forced=True))


def test_dependency_template_incompatible_blocks_filtered_export(tmp_path: Path) -> None:
    root = _with_ref_workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    _write_rows(root, "ItemType", [(1,)])
    _write_rows(root, "Item", [(1, 1)])
    # 依赖表 schema 改了但模板没更新
    write_yaml(
        root / "config" / "schemas" / "ItemType.yaml",
        {
            "table": "ItemType",
            "primary": "Id",
            "fields": [
                {"name": "Id", "type": "int32"},
                {"name": "Label", "type": "string", "comment": "标签"},
            ],
        },
    )
    from ct.app.exporting.models import ExportRequest
    from ct.app.exporting.service import run_export

    from ct.app.canonical_commands import CanonicalValidationError

    with pytest.raises(CanonicalValidationError) as excinfo:
        run_export(ExportRequest(root=root, table_filter="Item", forced=True))
    assert any("ItemType" in issue.table for issue in excinfo.value.issues)


def test_extra_trailing_columns_stay_a_warning(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    _write_rows(root, "Item", [(1, "x", 2)])
    path = root / "excel" / "Item.xlsx"
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2, column=4, value="Note\nstring")  # 未受管尾列
    workbook.save(path)

    workspace = CanonicalWorkspace.load(root)
    table = next(item for item in workspace.tables if item.table == "Item")
    layout = build_layout(
        table,
        schema_hash=compute_schema_hash(table, (*workspace.records, *workspace.enums)),
        records={record.name: record for record in workspace.records},
    )
    manifest = load_manifest(workspace.resolve("excel_dir") / "layout_manifests", "Item")
    assert check_reading_compatibility("Item", layout, path, manifest=manifest).ok


# --------------------------------------------------------------------------- #
# 4.3 explicit template update owns data migration
# --------------------------------------------------------------------------- #


def test_gen_template_refuses_unprovable_rename_mapping(tmp_path: Path) -> None:
    root = _workspace(tmp_path)
    canonical_gen_template(root, all_tables=True)
    _write_rows(root, "Item", [(1, "kept", 2)])
    workbook_before = (root / "excel" / "Item.xlsx").read_bytes()
    manifest_before = (root / "excel" / "layout_manifests" / "Item.json").read_bytes()

    # YAML 里把 Name 改名 DisplayName，旧 Name 列仍有数据
    write_yaml(
        root / "config" / "schemas" / "Item.yaml",
        _item_schema(
            fields=[
                {"name": "Id", "type": "int32"},
                {"name": "DisplayName", "type": "string", "comment": "名称"},
                {"name": "Level", "type": "int32", "comment": "等级"},
            ]
        ),
    )
    with pytest.raises(Exception):
        canonical_gen_template(root, all_tables=True)

    assert (root / "excel" / "Item.xlsx").read_bytes() == workbook_before
    assert (root / "excel" / "layout_manifests" / "Item.json").read_bytes() == manifest_before
