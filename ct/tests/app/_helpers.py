"""Shared  canonical workspace builders for schema tests."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import yaml


def write_yaml(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )


def build_project(
    root: Path,
    *,
    schemas: list[dict[str, Any]] | None = None,
    types: list[dict[str, Any]] | None = None,
) -> Path:
    """Build a minimal  workspace (config + optional schemas/types)."""
    (root / "config").mkdir(parents=True, exist_ok=True)
    (root / "config" / "schemas").mkdir(parents=True, exist_ok=True)
    (root / "config" / "types").mkdir(parents=True, exist_ok=True)
    write_yaml(
        root / "config" / "global.yaml",
        {
            "primary_lang": "zh",
            "secondary_langs": ["en"],
        },
    )
    for schema in schemas or []:
        write_yaml(root / "config" / "schemas" / f"{schema['table']}.yaml", schema)
    for type_def in types or []:
        write_yaml(root / "config" / "types" / f"{type_def['name']}.yaml", type_def)
    return root


def make_workbook(root: Path, table: str, rows: list[list] | None = None) -> Path:
    """Generate the canonical template for ``table``, then write ``rows`` below it.

    validate/export verify the managed header structure before reading, so test
    workbooks must be real templates instead of hand-rolled header rows.
    """
    from openpyxl import load_workbook

    from ct.app.canonical_commands import canonical_gen_template
    from ct.excel.layout_manifest import load_manifest

    canonical_gen_template(root, table_filter=table)
    path = root / "excel" / f"{table}.xlsx"
    manifest = load_manifest(root / "excel" / "layout_manifests", table)
    start = (manifest.header_rows if manifest is not None else 2) + 1
    workbook = load_workbook(path)
    worksheet = workbook.active
    for offset, row in enumerate(rows if rows is not None else [[1, "剑", 10], [2, "盾", 20]]):
        for column, value in enumerate(row, start=1):
            worksheet.cell(row=start + offset, column=column, value=value)
    workbook.save(path)
    workbook.close()
    return path


def set_cell(root: Path, table: str, row_offset: int, column: int, value: object) -> None:
    """Write one data cell (``row_offset`` 1 = first data row) of a template."""
    from openpyxl import load_workbook

    from ct.excel.layout_manifest import load_manifest

    path = root / "excel" / f"{table}.xlsx"
    manifest = load_manifest(root / "excel" / "layout_manifests", table)
    start = (manifest.header_rows if manifest is not None else 2) + 1
    workbook = load_workbook(path)
    workbook.active.cell(row=start + row_offset - 1, column=column, value=value)
    workbook.save(path)
    workbook.close()
