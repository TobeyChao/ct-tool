"""Incremental exports must be byte-equivalent to forced exports."""
from pathlib import Path
import json

import pytest
from openpyxl import Workbook

from _helpers import build_project, write_yaml
from ct.app.canonical_export import run_canonical_export
from ct.app.exporting import build as export
from ct.app.canonical_commands import CanonicalValidationError


def _excel(root, name="Item", rows=None):
    path = root / "excel" / f"{name}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["Id", "Name", "Value"])
    ws.append(["主键", "名称", "值"])
    for row in rows if rows is not None else [[1, "剑", 10], [2, "盾", 20]]:
        ws.append(row)
    wb.save(path)
    wb.close()


@pytest.fixture
def project(tmp_path):
    root = build_project(tmp_path / "gd", schemas=[{
        "table": name, "primary": "Id", "fields": [
            {"name": "Id", "type": "int32"},
            {"name": "Name", "type": "string", "i18n": True},
            {"name": "Value", "type": "int32"},
        ],
    } for name in ("Item", "Other")])
    _excel(root)
    _excel(root, "Other")
    return root


def _snapshot(root):
    return {str(p.relative_to(root)): (p.read_bytes(), p.stat().st_mtime_ns)
            for p in (root / "output").rglob("*") if p.is_file()}


def _assert_forced_equal(root):
    before = _snapshot(root)
    result = run_canonical_export(root, forced=True)
    assert result["cache_hits"] == 0
    assert len(result["written"]) == len(before)
    assert {p: b for p, (b, _) in before.items()} == {
        p: b for p, (b, _) in _snapshot(root).items()}


def test_warm_export_skips_generators_and_preserves_mtimes(project, monkeypatch):
    run_canonical_export(project)
    before = _snapshot(project)
    manifest = project / "excel/layout_manifests/Item.json"
    manifest_before = (manifest.read_bytes(), manifest.stat().st_mtime_ns)
    def fail(*args, **kwargs):
        pytest.fail("unchanged generator should have been reused")
    for name in ("build_canonical_table_bytes", "build_canonical_bundle",
                 "generate_csharp_accessor", "generate_lua_accessor",
                 "table_fbs_text", "types_fbs_text", "serialize_table_json"):
        monkeypatch.setattr(export, name, fail)
    result = run_canonical_export(project)
    assert not result["written"]
    assert result["cache_misses"] == 0
    assert result["cache_hits"] > 0
    assert before == _snapshot(project)
    assert manifest_before == (manifest.read_bytes(), manifest.stat().st_mtime_ns)


def test_translation_only_rebuilds_affected_language(project):
    run_canonical_export(project)
    path = project / "i18n/en/Item.json"
    path.parent.mkdir(parents=True)
    entry = {"1.Name": {"source": "剑", "text": "Sword", "confirmed": True}}
    path.write_text(json.dumps(entry))
    result = run_canonical_export(project)
    assert {Path(p).relative_to(project).as_posix() for p in result["written"]} == {
        "output/json/Item_en.json", "output/binary/data_en.bin"}
    entry["1.Name"]["status"] = "translated"
    entry["999.Name"] = {"text": "orphan", "confirmed": True}
    path.write_text(json.dumps(entry, indent=4))
    assert run_canonical_export(project)["cache_misses"] == 0
    _assert_forced_equal(project)


def test_data_edit_keeps_other_table_and_schema_outputs(project):
    run_canonical_export(project)
    before = _snapshot(project)
    _excel(project, rows=[[1, "新剑", 11], [2, "盾", 20]])
    result = run_canonical_export(project)
    assert {Path(p).name for p in result["written"]} == {
        "Item_zh.json", "Item_en.json", "data_zh.bin", "data_en.bin"}
    for path, value in _snapshot(project).items():
        if "Other" in path or "/fbs/" in path or "/generated/" in path:
            assert value == before[path]
    _assert_forced_equal(project)


def test_uniform_transition_invalidates_accessor(project):
    run_canonical_export(project)
    manifest = project / "excel/layout_manifests/Item.json"
    assert json.loads(manifest.read_text())["uniform"] is True
    _excel(project, rows=[[1, "剑", 0], [2, "盾", 0]])
    result = run_canonical_export(project)
    assert json.loads(manifest.read_text())["uniform"] is False
    assert "ItemAccessor.cs" in {Path(p).name for p in result["written"]}
    _assert_forced_equal(project)


def test_missing_and_modified_outputs_are_repaired(project):
    run_canonical_export(project)
    before = _snapshot(project)
    (project / "output/binary/data_zh.bin").unlink()
    (project / "output/generated/csharp/ItemAccessor.cs").write_text("broken")
    result = run_canonical_export(project)
    assert len(result["written"]) == 2
    assert result["cache_misses"] == 0
    assert {p: b for p, (b, _) in before.items()} == {
        p: b for p, (b, _) in _snapshot(project).items()}


@pytest.mark.parametrize("corrupt", ["null", "{", '{"payload":"!"}'])
def test_corrupt_cache_fails_safe(project, corrupt):
    run_canonical_export(project)
    before = _snapshot(project)
    for path in (project / "cache/artifacts").rglob("*.json"):
        path.write_text(corrupt)
    assert run_canonical_export(project)["cache_misses"] > 0
    assert before == _snapshot(project)


def test_warm_cache_does_not_skip_validation(project):
    run_canonical_export(project)
    before = _snapshot(project)
    _excel(project, rows=[[1, "剑", 10], [1, "盾", 20]])
    with pytest.raises(CanonicalValidationError):
        run_canonical_export(project)
    assert before == _snapshot(project)


def test_version_change_invalidates_generators(project, monkeypatch):
    run_canonical_export(project)
    monkeypatch.setattr(export, "CODEGEN_VERSION", "next")
    result = run_canonical_export(project)
    assert result["cache_misses"] > 0
    assert not result["written"]


def test_full_export_after_filtered_export_restores_complete_bundle(project):
    run_canonical_export(project, table_filter="Item", lang_filter="en")
    run_canonical_export(project)
    _assert_forced_equal(project)


def test_removed_table_and_language_cleanup(project):
    run_canonical_export(project)
    (project / "config/schemas/Other.yaml").unlink()
    write_yaml(project / "config/global.yaml", {"primary_lang": "zh", "secondary_langs": []})
    run_canonical_export(project)
    paths = _snapshot(project)
    assert not any("Other" in p or "_en." in p for p in paths)
    _assert_forced_equal(project)


def test_transitive_record_change_rebuilds_dependent_binary_and_accessor(tmp_path):
    from openpyxl import load_workbook
    from ct.app.canonical_commands import canonical_gen_template
    root = build_project(tmp_path / "gd", schemas=[{
        "table": "Item", "primary": "Id", "fields": [
            {"name": "Id", "type": "int32"},
            {"name": "Attributes", "type": "Stats"},
        ],
    }], types=[
        {"kind": "record", "name": "Stats", "fields": [{"name": "Strength", "type": "Power"}]},
        {"kind": "record", "name": "Power", "fields": [{"name": "Amount", "type": "int32"}]},
    ])
    canonical_gen_template(root, all_tables=True)
    path = root / "excel/Item.xlsx"
    wb = load_workbook(path)
    wb.active.cell(7, 1, 1)
    wb.active.cell(7, 2, 10)
    wb.save(path)
    wb.close()
    run_canonical_export(root)
    write_yaml(root / "config/types/Power.yaml", {
        "kind": "record", "name": "Power", "fields": [{"name": "Amount", "type": "int64"}],
    })
    result = run_canonical_export(root)
    assert {"data_zh.bin", "ItemAccessor.cs", "types.fbs"} <= {
        Path(p).name for p in result["written"]}
    _assert_forced_equal(root)


def test_unrelated_named_type_does_not_rebuild_table_generators(project, monkeypatch):
    run_canonical_export(project)
    write_yaml(project / "config/types/Unused.yaml", {
        "kind": "record", "name": "Unused", "fields": [{"name": "Amount", "type": "int64"}],
    })
    def fail(*args, **kwargs):
        pytest.fail("unrelated named resource must not invalidate this table")
    monkeypatch.setattr(export, "build_canonical_table_bytes", fail)
    monkeypatch.setattr(export, "generate_csharp_accessor", fail)
    result = run_canonical_export(project)
    assert {Path(p).name for p in result["written"]} == {"types.fbs"}


def test_full_export_reclaims_obsolete_cache_entries(project):
    run_canonical_export(project)
    old_entries = set((project / "cache/artifacts").glob("*/*.json"))
    _excel(project, rows=[[1, "新剑", 100], [2, "盾", 20]])
    run_canonical_export(project, table_filter="Item")
    assert all(p.exists() for p in old_entries)  # partial exports retain other scopes
    run_canonical_export(project)
    assert any(not p.exists() for p in old_entries)
    assert run_canonical_export(project)["cache_misses"] == 0


def test_checksum_mismatch_rebuilds_cached_payload(project):
    import base64
    run_canonical_export(project)
    before = _snapshot(project)
    for path in (project / "cache/artifacts/build_canonical_bundle").glob("*.json"):
        entry = json.loads(path.read_text())
        entry["payload"] = base64.b64encode(b"corrupt").decode("ascii")
        path.write_text(json.dumps(entry))
    result = run_canonical_export(project)
    assert result["cache_misses"] == 2
    assert before == _snapshot(project)


def test_unrelated_types_preserve_manifest_and_status(project):
    from ct.app.canonical_commands import canonical_status
    run_canonical_export(project)
    manifest = project / "excel/layout_manifests/Item.json"
    before = (manifest.read_bytes(), manifest.stat().st_mtime_ns)
    for field_type in ("int32", "int64"):
        write_yaml(project / "config/types/Unused.yaml", {
            "kind": "record", "name": "Unused", "fields": [{"name": "Amount", "type": field_type}],
        })
        write_yaml(project / "config/types/UnusedEnum.yaml", {
            "kind": "enum", "name": "UnusedEnum", "values": [{"name": "Common", "comment": field_type}],
        })
        assert canonical_status(project)["drifted"] == []
        run_canonical_export(project)
        assert before == (manifest.read_bytes(), manifest.stat().st_mtime_ns)
        assert canonical_status(project)["drifted"] == []
    (project / "config/types/Unused.yaml").unlink()
    (project / "config/types/UnusedEnum.yaml").unlink()
    assert canonical_status(project)["drifted"] == []
    run_canonical_export(project)
    assert before == (manifest.read_bytes(), manifest.stat().st_mtime_ns)


def test_referenced_enum_drift_and_template_export_hash_agree(tmp_path):
    from ct.app.canonical_commands import canonical_gen_template, canonical_status
    root = build_project(tmp_path / "gd", schemas=[{
        "table": "Item", "primary": "Id", "fields": [
            {"name": "Id", "type": "int32"},
            {"name": "Rarity", "type": "ItemRarity"},
        ],
    }], types=[{
        "kind": "enum", "name": "ItemRarity", "values": [{"name": "Common", "comment": "普通"}],
    }])
    canonical_gen_template(root, all_tables=True)
    manifest = root / "excel/layout_manifests/Item.json"
    first_hash = json.loads(manifest.read_text())["schema_hash"]
    assert canonical_status(root)["drifted"] == []
    run_canonical_export(root)
    assert json.loads(manifest.read_text())["schema_hash"] == first_hash
    write_yaml(root / "config/types/ItemRarity.yaml", {
        "kind": "enum", "name": "ItemRarity", "values": [
            {"name": "Common", "comment": "普通"}, {"name": "Rare", "comment": "稀有"},
        ],
    })
    assert canonical_status(root)["drifted"] == ["Item"]
    canonical_gen_template(root, all_tables=True)
    second_hash = json.loads(manifest.read_text())["schema_hash"]
    assert second_hash != first_hash
    assert canonical_status(root)["drifted"] == []
    run_canonical_export(root)
    assert json.loads(manifest.read_text())["schema_hash"] == second_hash
