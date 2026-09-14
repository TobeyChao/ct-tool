"""End-to-end workflow: edit -> save YAML -> template gate -> explicit template
update -> export (task 5.3)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ct.app.canonical_commands import CanonicalValidationError, canonical_gen_template
from ct.app.canonical_workspace import CanonicalWorkspace
from ct.app.exporting.models import ExportRequest
from ct.app.exporting.service import run_export
from ct.app.schema_workspace.candidate import candidate_hash
from ct.app.schema_workspace.commands_reducer import Command, DraftLog
from ct.web.app import create_app

from _helpers import build_project


def _project(tmp_path: Path) -> Path:
    return build_project(
        tmp_path / "gd",
        schemas=[
            {
                "table": "Item",
                "primary": "Id",
                "fields": [
                    {"name": "Id", "type": "int32"},
                    {"name": "Name", "type": "string", "i18n": True, "comment": "名称"},
                ],
            }
        ],
    )


def _candidate_hash(root: Path, commands: list[dict]) -> str:
    workspace = CanonicalWorkspace.load(root)
    log = DraftLog(workspace.resources.resources, base_indexes=dict(workspace.indexes))
    for item in commands:
        log.execute(Command(item["type"], item["payload"]))
    resources, indexes = log.current()
    return candidate_hash(resources, indexes)


def test_edit_save_template_gate_and_export(tmp_path: Path) -> None:
    root = _project(tmp_path)
    canonical_gen_template(root, all_tables=True)
    # 模板 + 一行数据，先跑通一次成功导出
    workbook_path = root / "excel" / "Item.xlsx"
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    sheet.cell(row=3, column=1, value=1)
    sheet.cell(row=3, column=2, value="铁剑")
    workbook.save(workbook_path)
    run_export(ExportRequest(root=root, forced=True))

    client = create_app(root).test_client()
    revision = client.get("/api/schema-workspace").get_json()["data"]["schemaRevision"]
    commands = [
        {
            "type": "add_field",
            "payload": {"owner": "table:Item", "field": {"name": "Price", "type": "int32"}},
        }
    ]
    candidate = _candidate_hash(root, commands)

    # 1) 保存只写 YAML：Excel、产物、账本都不动
    xlsx_before = workbook_path.read_bytes()
    artifacts_before = {
        str(path.relative_to(root)): path.read_bytes()
        for path in (root / "output").rglob("*")
        if path.is_file()
    }
    ledger = root / "cache" / "state.json"
    ledger_before = ledger.read_bytes()

    resp = client.post(
        "/api/schema-workspace/save",
        json={"schemaRevision": revision, "commands": commands, "candidateHash": candidate},
    )
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()["data"]
    assert [Path(path).name for path in data["written"]] == ["Item.yaml"]
    assert "Price" in (root / "config" / "schemas" / "Item.yaml").read_text(encoding="utf-8")
    assert workbook_path.read_bytes() == xlsx_before
    assert {
        str(path.relative_to(root)): path.read_bytes()
        for path in (root / "output").rglob("*")
        if path.is_file()
    } == artifacts_before
    assert ledger.read_bytes() == ledger_before

    # 2) 模板没更新：导出被读取闸门拒绝，且不刷新 manifest / 产物 / 账本
    manifests_before = {
        path.name: path.read_bytes()
        for path in (root / "excel" / "layout_manifests").glob("*.json")
    }
    with pytest.raises(CanonicalValidationError) as excinfo:
        run_export(ExportRequest(root=root, forced=True))
    assert any(issue.table == "Item" for issue in excinfo.value.issues)
    assert {
        path.name: path.read_bytes()
        for path in (root / "excel" / "layout_manifests").glob("*.json")
    } == manifests_before
    assert ledger.read_bytes() == ledger_before

    # 3) 显式更新模板（保留旧数据），填数据后导出成功
    messages = canonical_gen_template(root, all_tables=True)
    assert messages == ["模板已生成: Item"]
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    sheet.cell(row=3, column=3, value=99)
    workbook.save(workbook_path)

    result = run_export(ExportRequest(root=root, forced=True))
    assert result is not None
    payload = json.loads((root / "output" / "json" / "Item_zh.json").read_text(encoding="utf-8"))
    assert payload["Items"][0]["Price"] == 99
    assert payload["Items"][0]["Name"] == "铁剑"
    assert ledger.read_bytes() != ledger_before


def test_save_without_excel_or_manifest_still_succeeds(tmp_path: Path) -> None:
    """保存只改 YAML：没有 Excel / manifest 也不阻塞，导出才需要模板。"""
    root = _project(tmp_path)
    client = create_app(root).test_client()
    revision = client.get("/api/schema-workspace").get_json()["data"]["schemaRevision"]
    commands = [
        {
            "type": "set_property",
            "payload": {"owner": "table:Item", "name": "Name", "property": "comment", "value": "改"},
        }
    ]
    candidate = _candidate_hash(root, commands)
    resp = client.post(
        "/api/schema-workspace/save",
        json={"schemaRevision": revision, "commands": commands, "candidateHash": candidate},
    )
    assert resp.status_code == 200
    assert resp.get_json()["data"]["isNoOp"] is False
