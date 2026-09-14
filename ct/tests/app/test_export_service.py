"""统一完成服务：持锁恢复、完成策略、部署失败与账本语义（task 5.2）。"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from _helpers import build_project
from ct.app.exporting.models import CompletionPolicy, ExportRequest
from ct.app.exporting.service import run_export
from ct.storage.workspace_lock import WorkspaceBusyError, WorkspaceLock


def _excel_bytes(root: Path, ids: list[int]) -> bytes:
    """该工程模板 + 指定主键的字节（读取闸门要求真实受管表头）。"""
    from openpyxl import load_workbook

    from ct.excel.layout_manifest import load_manifest

    path = root / "excel" / "Item.xlsx"
    manifest = load_manifest(root / "excel" / "layout_manifests", "Item")
    start = (manifest.header_rows if manifest is not None else 2) + 1
    workbook = load_workbook(path)
    sheet = workbook.active
    for index in range(start, sheet.max_row + 1):
        sheet.cell(row=index, column=1, value=None)
        sheet.cell(row=index, column=2, value=None)
    for offset, value in enumerate(ids):
        sheet.cell(row=start + offset, column=1, value=value)
        sheet.cell(row=start + offset, column=2, value="剑")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _project(tmp_path: Path, tables: tuple[str, ...] = ("Item",)) -> Path:
    root = build_project(
        tmp_path / "gd",
        schemas=[
            {
                "table": name,
                "primary": "Id",
                "fields": [
                    {"name": "Id", "type": "int32"},
                    {"name": "Name", "type": "string", "i18n": True},
                ],
            }
            for name in tables
        ],
    )
    from _helpers import make_workbook

    for name in tables:
        make_workbook(root, name, [[1, "剑"]])
    return root


def _snapshot(root: Path) -> dict[str, bytes]:
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in (root / "output").rglob("*")
        if path.is_file()
    }


def _ledger(root: Path) -> bytes:
    return (root / "cache" / "state.json").read_bytes()


def test_export_only_strategy_publishes_and_records_the_ledger(
    tmp_path: Path,
) -> None:
    root = _project(tmp_path)
    result = run_export(ExportRequest(root=root))

    assert result.tables == 1
    assert (root / "output" / "json" / "Item_zh.json").is_file()
    assert (root / "cache" / "state.json").is_file()
    assert not (root / ".ct" / "export-publication.json").exists()


def test_service_takes_the_lock_exactly_once(tmp_path: Path) -> None:
    """进程内锁不可重入：若内部递归抢锁会立刻 busy，能跑通即证明没有递归。"""
    root = _project(tmp_path)
    assert run_export(ExportRequest(root=root)).tables == 1


def test_busy_workspace_is_reported(tmp_path: Path) -> None:
    root = _project(tmp_path)
    with WorkspaceLock(root):
        with pytest.raises(WorkspaceBusyError):
            run_export(ExportRequest(root=root))


def test_deploy_failure_keeps_new_local_artifacts_and_old_ledger(
    tmp_path: Path, monkeypatch
) -> None:
    root = _project(tmp_path)
    run_export(ExportRequest(root=root))
    ledger_before = _ledger(root)
    local_before = _snapshot(root)

    def boom(*args, **kwargs):
        raise OSError("注入：Unity 目标不可写")

    monkeypatch.setattr("ct.export.deploy.deploy", boom)

    # 数据变了，本次强制导出会产出**新**本地产物
    (root / "excel" / "Item.xlsx").write_bytes(_excel_bytes(root, [2]))

    with pytest.raises(OSError):
        run_export(
            ExportRequest(root=root, forced=True),
            policy=CompletionPolicy.export_then_deploy(),
        )

    # 本地已完整发布新版本；账本停留在上一次成功记录
    assert _snapshot(root) != local_before
    payload = json.loads(
        (root / "output" / "json" / "Item_zh.json").read_text("utf-8")
    )
    assert [row["Id"] for row in payload["Items"]] == [2]
    assert _ledger(root) == ledger_before


def test_ledger_failure_is_not_reported_as_success(
    tmp_path: Path, monkeypatch
) -> None:
    root = _project(tmp_path)

    def boom(*args, **kwargs):
        raise OSError("注入：账本替换失败")

    monkeypatch.setattr("ct.app.exporting.build.persist_export_state", boom)
    ledger = root / "cache" / "state.json"
    ledger_before = ledger.read_bytes() if ledger.exists() else None

    with pytest.raises(OSError):
        run_export(ExportRequest(root=root))

    # 已完整发布的本地产物保留，但成功账本不得被推进
    assert (root / "output" / "json" / "Item_zh.json").is_file()
    assert (ledger.read_bytes() if ledger.exists() else None) == ledger_before


def test_partial_scope_keeps_other_ledger_records(tmp_path: Path) -> None:
    root = _project(tmp_path, tables=("Item", "Other"))
    run_export(ExportRequest(root=root))
    full = json.loads(_ledger(root).decode("utf-8"))
    assert {"Item", "Other"} <= set(full["excel_hashes"])
    assert set(full["bundles"]) == {"zh", "en"}

    run_export(ExportRequest(root=root, table_filter="Item", forced=True))
    after = json.loads(_ledger(root).decode("utf-8"))

    assert after["excel_hashes"]["Other"] == full["excel_hashes"]["Other"]
    assert set(after["bundles"]) == set(full["bundles"])


def test_notify_fires_after_local_publication_and_before_deploy(
    tmp_path: Path, monkeypatch
) -> None:
    root = _project(tmp_path)
    order: list[str] = []

    def stub_deploy(config, for_build, reporter):
        order.append("deploy")
        return 0

    monkeypatch.setattr("ct.export.deploy.deploy", stub_deploy)

    def notify(result):
        order.append("notify")
        # 通知时本地发布必须已经完成
        assert (root / "output" / "json" / "Item_zh.json").is_file()

    run_export(
        ExportRequest(root=root),
        policy=CompletionPolicy.export_then_deploy(for_build=True),
        notify=notify,
    )

    assert order == ["notify", "deploy"]


def test_web_style_export_only_never_deploys(tmp_path: Path, monkeypatch) -> None:
    """Web 策略（export_only）即使配置启用了 deploy 也不部署。"""
    root = _project(tmp_path)
    (root / "config" / "global.yaml").write_text(
        "primary_lang: zh\nsecondary_langs: [en]\n"
        "deploy:\n  enabled: true\n  unity_project: unity\n"
        "  targets:\n    - {source: output/json, dest: Assets/Config}\n",
        encoding="utf-8",
    )
    calls: list[str] = []
    monkeypatch.setattr(
        "ct.export.deploy.deploy",
        lambda *args, **kwargs: calls.append("deploy") or 0,
    )

    run_export(ExportRequest(root=root), policy=CompletionPolicy.export_only())

    assert calls == [], "export_only 策略不得部署"
    assert (root / "cache" / "state.json").is_file()


def test_export_recovers_pending_publication_before_running(tmp_path: Path) -> None:
    """服务入口必须先恢复未完成的发布，再执行本次请求。"""
    root = _project(tmp_path)
    run_export(ExportRequest(root=root))

    publisher_dir = root / ".ct"
    publisher_dir.mkdir(parents=True, exist_ok=True)
    target = root / "output" / "json" / "Item_zh.json"
    backup = publisher_dir / "backup" / "op" / "Item_zh.json"
    backup.parent.mkdir(parents=True, exist_ok=True)
    backup.write_bytes(target.read_bytes())
    target.write_bytes(b'{"Items": "broken"}')
    (publisher_dir / "export-publication.json").write_text(
        json.dumps(
            {
                "format": "export-publication/1",
                "operation_id": "op",
                "root": str(root.resolve()),
                "phase": "publishing",
                "allowed_dirs": [str((root / "output").resolve())],
                "entries": [
                    {
                        "path": str(target.resolve()),
                        "op": "replace",
                        "existed": True,
                        "old_hash": "x",
                        "new_hash": "y",
                        "staged": None,
                        "backup": str(backup.resolve()),
                        "done": True,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    run_export(ExportRequest(root=root, forced=True))

    assert target.read_bytes() != b'{"Items": "broken"}'
    assert not (publisher_dir / "export-publication.json").exists()
