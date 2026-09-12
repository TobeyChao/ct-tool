"""Canonical CLI command implementations (validate/status/gen-template/i18n).

These back the canonical-only CLI and Web; the legacy (pre) path is removed.
"""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from ct.app.canonical_workspace import CanonicalWorkspace
from ct.app.data_preparation import prepare_tables, records_map
from ct.cache.canonical_state import (
    CanonicalCacheState,
    load_state,
    record_excel_hashes,
    save_state,
)
from ct.diagnostics.errors import Issue, IssueCode, WorkspaceIssue
from ct.excel.canonical_reader import read_canonical_excel
from ct.excel.canonical_template import generate_canonical_template
from ct.excel.layout import Column, Layout, build_layout
from ct.excel.layout_manifest import LayoutManifest, save_manifest
from ct.excel.planning import plan_excel_migration
from ct.schema.hashing import compute_schema_hash
from ct.storage.publication import FilePublisher, PublicationError


class CanonicalValidationError(ValueError):
    """Canonical 校验失败：携带结构化 issues，供 CLI / Web 渲染。"""

    def __init__(self, issues: list[Issue]) -> None:
        super().__init__(f"校验发现 {len(issues)} 个问题")
        self.issues = issues


def canonical_publication_state(root: Path) -> str | None:
    """未完成或损坏的发布恢复记录描述；没有则返回 ``None``。

    **只读**：不执行恢复、不写任何文件。损坏记录返回错误描述而不是抛错，
    让 validate/status 能把它当作一个问题报出来（而不是静默报告「正常」）。
    """
    publisher = FilePublisher(root)
    try:
        journal = publisher.read_journal()
    except PublicationError as exc:
        return str(exc)
    if journal is None:
        return None
    return (
        f"存在未完成的发布（operation {journal.operation_id}，阶段 {journal.phase}）"
        f"——下一次 export/deploy 会先恢复，或人工检查 {publisher.journal_path}"
    )


def canonical_validate(
    root: Path,
    *,
    table_filter: str | None = None,
) -> list[Issue]:
    """Read + validate a canonical workspace; returns structured issues.

    Full validation: Excel read (type coercion), primary empty/duplicate,
    CodeName index data gate and cross-table ``ref`` foreign-key values (must
    exist in the target table's primary-key set). The reads and the validation
    rules come from the shared kernel in :mod:`ct.app.data_preparation`; this
    wrapper keeps only the ``list[Issue]`` shape and the unknown-table
    diagnostic. The legacy path no longer exists.

    未完成/损坏的发布记录也会作为工作区级问题报出（只读检测，不做恢复），
    避免校验通过的同时 export 因未恢复的现场被拒。
    """
    workspace = CanonicalWorkspace.load(root)
    result = prepare_tables(workspace, table_filter=table_filter)
    if result.unknown_table is not None:
        return [
            WorkspaceIssue(
                "", IssueCode.WORKSPACE, f"表 '{result.unknown_table}' 不存在"
            )
        ]
    issues = list(result.issues)
    publication = canonical_publication_state(root)
    if publication is not None:
        issues.append(WorkspaceIssue("", IssueCode.WORKSPACE, publication))
    return issues


def _file_sha256(path: Path) -> str:
    if not path.exists():
        return ""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical_status(root: Path) -> dict[str, list[str]]:
    """Per-table data-change + template-drift status for a canonical workspace.

    `changed` reports a real data edit pending export: the current Excel file
    hash differs from the `excel_hashes` ledger recorded at the last
    export/gen-template, or the table has never been exported (no cache entry).
    `drifted` reports template/schema drift: the layout manifest's
    `schema_hash` no longer matches the current schema.
    """
    ws = CanonicalWorkspace.load(root)
    records = records_map(ws)
    excel_dir = ws.resolve("excel_dir")
    cache_dir = ws.resolve("cache_dir")
    manifest_dir = ws.resolve("excel_dir") / "layout_manifests"
    state = load_state(cache_dir)
    changed: list[str] = []
    drifted: list[str] = []
    missing: list[str] = []
    for table in ws.tables:
        excel_path = excel_dir / (table.excel_file or f"{table.table}.xlsx")
        if not excel_path.exists():
            missing.append(table.table)
            continue
        current_hash = _file_sha256(excel_path)
        manifest = _load_manifest(manifest_dir, table.table)
        schema_hash = compute_schema_hash(table, (*ws.records, *ws.enums))
        layout = build_layout(table, schema_hash=schema_hash, records=records)
        workbook_column_count = _template_column_count(excel_path)
        if (
            manifest is None
            or manifest.schema_hash != schema_hash
            or workbook_column_count != layout.column_count
        ):
            drifted.append(table.table)
        cached_hash = state.excel_hashes.get(table.table) if state else None
        if cached_hash is None or cached_hash != current_hash:
            changed.append(table.table)
    return {"changed": sorted(set(changed)), "drifted": sorted(set(drifted)), "missing": sorted(missing)}


def _load_manifest(manifest_dir: Path, table: str) -> LayoutManifest | None:
    from ct.excel.layout_manifest import load_manifest

    return load_manifest(manifest_dir, table)


def _layout_from_manifest(table_id: str, manifest: LayoutManifest) -> Layout:
    """Rebuild the previous layout needed to migrate an existing workbook."""
    columns = tuple(
        Column(
            index=int(item["index"]),
            stable_path=str(item["stablePath"]),
            type_text=str(item.get("typeExpr", "")),
            annotation=str(item.get("annotation", "")),
            leaf=str(item.get("leaf", "")),
            group_index=(
                int(item["groupIndex"]) if item.get("groupIndex") is not None else None
            ),
            depth=int(item.get("depth", 1)),
        )
        for item in manifest.columns
    )
    return Layout(
        table_id=table_id,
        schema_hash=manifest.schema_hash,
        header_rows=manifest.header_rows,
        columns=columns,
    )


def _migrate_excel_rows(
    old_path: Path,
    new_path: Path,
    old_layout: Layout,
    new_layout: Layout,
    manifest: LayoutManifest,
) -> None:
    """Copy old data rows into a newly generated workbook by stable column path."""
    plan = plan_excel_migration(
        old_layout,
        new_layout,
        old_path,
        manifest=manifest,
    )
    if plan.blocked:
        details = "；".join(issue.render() for issue in plan.issues)
        raise ValueError(f"Excel 数据无法安全迁移：{details}")

    targets = {
        migration.old_index: migration.new_index
        for migration in plan.migrations
        if migration.new_index is not None
    }
    old_wb = load_workbook(str(old_path), read_only=True, data_only=False)
    # Preserve CellRichText header runs while saving migrated data; loading
    # without rich_text=True permanently flattens them to plain strings.
    new_wb = load_workbook(
        str(new_path), read_only=False, data_only=False, rich_text=True
    )
    try:
        old_ws = old_wb.active
        new_ws = new_wb.active
        new_row = new_layout.header_rows + 1
        for row_index in range(old_layout.header_rows + 1, old_ws.max_row + 1):
            values = {
                old_index: old_ws.cell(row=row_index, column=old_index).value
                for old_index in targets
            }
            if not any(value is not None and (not isinstance(value, str) or value.strip()) for value in values.values()):
                continue
            for old_index, new_index in targets.items():
                new_ws.cell(row=new_row, column=new_index).value = values[old_index]
            new_row += 1
        new_wb.save(str(new_path))
    finally:
        old_wb.close()
        new_wb.close()


def _template_column_count(path: Path) -> int | None:
    """Read the actual number of columns in the Excel template."""
    if not path.exists():
        return None
    workbook = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=False)
        ws = workbook.active
        if ws is None:
            return None
        return int(ws.max_column)
    except (OSError, KeyError, TypeError, ValueError, AttributeError):
        return None
    finally:
        if workbook is not None:
            workbook.close()



def _validate_staged_workbook(path: Path) -> None:
    """Reject an incomplete or unreadable XLSX before publishing it."""
    try:
        with zipfile.ZipFile(path) as archive:
            broken_member = archive.testzip()
        if broken_member is not None:
            raise ValueError(f"XLSX ZIP 成员损坏: {broken_member}")

        workbook = load_workbook(str(path), read_only=True, data_only=False)
        try:
            if not workbook.sheetnames:
                raise ValueError("XLSX 中没有工作表")
        finally:
            workbook.close()
    except (OSError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ValueError(f"生成的 Excel 候选文件校验失败: {exc}") from exc


def _publish_staged_workbook(staged_path: Path, out_path: Path) -> None:
    """Atomically publish a staged workbook while preserving Windows metadata."""
    if os.name != "nt" or not out_path.exists():
        os.replace(staged_path, out_path)
        return

    # ReplaceFileW merges the replaced file's DACL and other metadata into the
    # replacement.  os.replace() only renames the replacement file on Windows,
    # so its ACL can unexpectedly become the destination ACL.
    import ctypes
    from ctypes import wintypes

    replace_file = ctypes.WinDLL("kernel32", use_last_error=True).ReplaceFileW
    replace_file.argtypes = [
        wintypes.LPCWSTR,
        wintypes.LPCWSTR,
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.LPVOID,
    ]
    replace_file.restype = wintypes.BOOL
    if not replace_file(
        str(out_path.resolve()),
        str(staged_path.resolve()),
        None,
        0,
        None,
        None,
    ):
        error = ctypes.get_last_error()
        raise OSError(error, ctypes.FormatError(error), str(out_path))


class UnknownTableError(ValueError):
    """表名不存在。

    CLI 与其他 `ValueError` 一样转为友好提示 + 退出码 1；Web 端点据类型映射为 404
    （`/api/schema-workspace/gen-template` 既有契约是 404 + `未找到表`）。
    """


def _require_table(ws: CanonicalWorkspace, name: str) -> str:
    """校验表名存在（精确匹配 PascalCase），返回原名。

    与 `canonical_validate` / `run_canonical_export` 同口径：未知表名一律报错，
    绝不静默处理 0 张表。仅大小写不符时提示正确写法——规格要求表名精确匹配
    PascalCase，这是最常见的输入错误。
    """
    names = [t.table for t in ws.tables]
    if name in names:
        return name
    hint = next((n for n in names if n.lower() == name.lower()), None)
    if hint is not None:
        raise UnknownTableError(
            f"表 '{name}' 不存在（是否想用 '{hint}'？表名精确匹配 PascalCase）"
        )
    raise UnknownTableError(f"表 '{name}' 不存在")


def _require_lang(ws: CanonicalWorkspace, lang: str) -> str:
    """校验语言在 `secondary_langs` 中，返回原名。"""
    langs = list(ws.config.secondary_langs)
    if lang in langs:
        return lang
    raise ValueError(
        f"语言 '{lang}' 不在 secondary_langs 中"
        f"（可用: {', '.join(langs) if langs else '无'}）"
    )


def canonical_gen_template(
    root: Path,
    *,
    table_filter: str | None = None,
    all_tables: bool = False,
) -> list[str]:
    """Generate canonical Excel templates + layout manifests."""
    ws = CanonicalWorkspace.load(root)
    records = records_map(ws)
    excel_dir = ws.resolve("excel_dir")
    excel_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = ws.resolve("cache_dir")
    manifest_dir = ws.resolve("excel_dir") / "layout_manifests"
    if table_filter is None and not all_tables:
        raise ValueError("请指定 --all 或 --table <表名>")
    if table_filter is not None:
        _require_table(ws, table_filter)
    targets = [t for t in ws.tables if table_filter is None or table_filter == t.table]
    messages: list[str] = []
    for table in targets:
        layout = build_layout(
            table,
            schema_hash=compute_schema_hash(table, (*ws.records, *ws.enums)),
            records=records,
        )
        out_path = excel_dir / (table.excel_file or f"{table.table}.xlsx")
        old_manifest = _load_manifest(manifest_dir, table.table)
        if out_path.exists() and old_manifest is not None:
            old_layout = _layout_from_manifest(table.resource_id, old_manifest)
            # Create the candidate directly under excel_dir so it receives the
            # same inherited Windows ACL and remains on the same volume.
            staged_fd, staged_name = tempfile.mkstemp(
                prefix=f".{out_path.stem}.",
                suffix=".staged.xlsx",
                dir=str(excel_dir),
            )
            os.close(staged_fd)
            staged_path = Path(staged_name)
            try:
                generate_canonical_template(
                    layout,
                    staged_path,
                    enums={e.name: e for e in ws.enums},
                    primary=table.primary,
                )
                _migrate_excel_rows(
                    out_path,
                    staged_path,
                    old_layout,
                    layout,
                    old_manifest,
                )
                _validate_staged_workbook(staged_path)
                _publish_staged_workbook(staged_path, out_path)
            finally:
                staged_path.unlink(missing_ok=True)
            save_manifest(
                manifest_dir,
                table.table,
                LayoutManifest.from_layout(layout),
            )
        else:
            if out_path.exists() and old_manifest is None:
                raise ValueError(
                    f"{table.table} 的 Excel 缺少布局 manifest，无法安全迁移；"
                    "请先备份后删除旧文件，再重新生成空模板"
                )
            generate_canonical_template(
                layout, out_path, enums={e.name: e for e in ws.enums}, primary=table.primary
            )
            save_manifest(manifest_dir, table.table, LayoutManifest.from_layout(layout))
        messages.append(f"模板已生成: {table.table}")
    if targets:
        state = load_state(cache_dir) or CanonicalCacheState()
        hashes = {
            t.table: _file_sha256(excel_dir / (t.excel_file or f"{t.table}.xlsx"))
            for t in targets
        }
        save_state(cache_dir, record_excel_hashes(state, hashes))
    return messages


def _i18n_progress(counts: dict[str, int]) -> float:
    """进度 = translated / (total - orphan)，无活跃条目视为 100%。"""
    active = counts["total"] - counts["orphan"]
    if active <= 0:
        return 1.0
    return round(counts["translated"] / active, 4)


def canonical_i18n_status(root: Path) -> dict[str, dict]:
    """Per-language + per-table translation counts for a canonical workspace."""
    from ct.export.i18n.state import compute_status

    ws = CanonicalWorkspace.load(root)
    config = ws.config
    i18n_dir = config.resolve("i18n_dir")
    i18n_tables = [t for t in ws.tables if any(f.i18n for f in t.fields)]
    tables = [t.table for t in i18n_tables]
    result: dict[str, dict] = {}
    for lang in config.secondary_langs:
        lang_dir = i18n_dir / lang
        lang_counts = {"translated": 0, "missing": 0, "stale": 0, "orphan": 0, "total": 0}
        table_detail: dict[str, dict] = {}
        for table in tables:
            source_path = i18n_dir / "source" / f"{table}.json"
            if not source_path.exists():
                continue
            source = json.loads(source_path.read_text(encoding="utf-8"))
            lang_path = lang_dir / f"{table}.json"
            entries = (
                json.loads(lang_path.read_text(encoding="utf-8"))
                if lang_path.exists()
                else {}
            )
            counts = {"translated": 0, "missing": 0, "stale": 0, "orphan": 0, "total": 0}
            for key, source_text in source.items():
                entry = entries.get(key) or {}
                text = str(entry.get("text", ""))
                confirmed = bool(entry.get("confirmed", False))
                status = compute_status(text, confirmed, in_source=True).value
                counts[status if status in counts else "missing"] += 1
                counts["total"] += 1
            for key, entry in entries.items():
                if key not in source:
                    counts["orphan"] += 1
                    counts["total"] += 1
            table_detail[table] = {**counts, "progress": _i18n_progress(counts)}
            for stat in ("translated", "missing", "stale", "orphan", "total"):
                lang_counts[stat] += counts[stat]
        result[lang] = {
            **lang_counts,
            "progress": _i18n_progress(lang_counts),
            "tables": table_detail,
        }
    return result


def canonical_i18n_sync(
    root: Path,
    *,
    table_filter: str | None = None,
    lang_filter: str | None = None,
    verbose: bool = False,
) -> list[str]:
    """Refresh source files and lang skeletons for a canonical workspace.

    ``lang_filter`` 只限定 **lang 骨架** 的写入范围：source 仍按选中的表全量刷新
    （规格 `ct i18n sync --lang`）。返回逐表（verbose 时含逐文件）消息，**末条为汇总**。
    """
    from ct.export.i18n.merger import write_lang_file, write_source_file
    from ct.export.i18n.state import sync_lang_table

    ws = CanonicalWorkspace.load(root)
    records = records_map(ws)
    config = ws.config
    excel_dir = config.resolve("excel_dir")
    i18n_dir = config.resolve("i18n_dir")
    if table_filter is not None:
        _i18n_table(ws, table_filter)
    langs = list(config.secondary_langs)
    if lang_filter is not None:
        _require_lang(ws, lang_filter)
        langs = [lang_filter]
    tables = [
        t for t in ws.tables
        if (table_filter is None or t.table == table_filter) and any(f.i18n for f in t.fields)
    ]
    messages: list[str] = []
    totals = {"added": 0, "updated": 0, "stale": 0, "orphan": 0}
    processed = 0
    for table in tables:
        i18n_fields = [f for f in table.fields if f.i18n]
        field_order = [f.name for f in i18n_fields]
        excel_path = excel_dir / (table.excel_file or f"{table.table}.xlsx")
        if not excel_path.exists():
            continue
        layout = build_layout(
            table,
            schema_hash=compute_schema_hash(table, (*ws.records, *ws.enums)),
            records=records,
        )
        parsed = read_canonical_excel(excel_path, layout, table, records=records, enums={e.name: e for e in ws.enums})
        source: dict[str, str] = {}
        for row in parsed.rows:
            row_id = row.get(table.primary)
            for field in i18n_fields:
                source[f"{row_id}.{field.name}"] = str(row.get(field.name, ""))
        write_source_file(source, i18n_dir / "source" / f"{table.table}.json", field_order)
        messages.append(f"synced {table.table}")
        if verbose:
            messages.append(
                f"写入 i18n/source/{table.table}.json（{len(source)} 条 source）"
            )
        processed += 1
        for lang in langs:
            lang_path = i18n_dir / lang / f"{table.table}.json"
            existing = (
                json.loads(lang_path.read_text(encoding="utf-8"))
                if lang_path.exists()
                else {}
            )
            synced = sync_lang_table(source, existing)
            added = sum(1 for key in synced if key not in existing)
            updated = sum(
                1 for key, entry in synced.items()
                if key in existing and existing[key] != entry
            )
            stale = sum(1 for e in synced.values() if e["status"] == "stale")
            orphan = sum(1 for e in synced.values() if e["status"] == "orphan")
            write_lang_file(synced, lang_path, field_order)
            for stat, count in (
                ("added", added), ("updated", updated),
                ("stale", stale), ("orphan", orphan),
            ):
                totals[stat] += count
            if verbose:
                messages.append(
                    f"写入 i18n/{lang}/{table.table}.json"
                    f"（新增 {added}、更新 {updated}、stale {stale}、orphan {orphan}）"
                )
    messages.append(
        f"处理 {processed} 张表 × {len(langs)} 语言：新增 {totals['added']}、"
        f"更新 {totals['updated']}、stale {totals['stale']}、orphan {totals['orphan']}"
    )
    return messages


def canonical_i18n_compact(
    root: Path,
    *,
    table_filter: str | None = None,
    lang_filter: str | None = None,
    dry_run: bool = False,
) -> dict:
    """Remove orphan entries from lang files.

    dry_run=True 时不落盘，返回即将删除的 files 明细供预览。
    """
    from ct.export.i18n.merger import write_lang_file

    ws = CanonicalWorkspace.load(root)
    config = ws.config
    i18n_dir = config.resolve("i18n_dir")
    if table_filter is not None:
        _i18n_table(ws, table_filter)
    langs = list(config.secondary_langs)
    if lang_filter is not None:
        _require_lang(ws, lang_filter)
        langs = [lang_filter]
    removed = 0
    touched = 0
    files: list[dict] = []
    for table in ws.tables:
        if table_filter is not None and table.table != table_filter:
            continue
        if not any(f.i18n for f in table.fields):
            continue
        field_order = [f.name for f in table.fields if f.i18n]
        source_path = i18n_dir / "source" / f"{table.table}.json"
        if not source_path.exists():
            continue
        source = set(json.loads(source_path.read_text(encoding="utf-8")).keys())
        for lang in langs:
            lang_path = i18n_dir / lang / f"{table.table}.json"
            if not lang_path.exists():
                continue
            entries = json.loads(lang_path.read_text(encoding="utf-8"))
            orphans = sorted(k for k in entries if k not in source)
            if not orphans:
                continue
            removed += len(orphans)
            touched += 1
            files.append(
                {"lang": lang, "table": table.table, "removed_keys": orphans}
            )
            if dry_run:
                continue
            for key in orphans:
                entries.pop(key, None)
            write_lang_file(entries, lang_path, field_order)
    return {
        "dry_run": dry_run,
        "touched": touched,
        "total_removed": removed,
        "files": files,
    }


def canonical_i18n_tables(root: Path) -> list[dict]:
    """List all tables with their i18n field metadata (for the picker)."""
    ws = CanonicalWorkspace.load(root)
    return [
        {
            "table": t.table,
            "field_count": len(t.fields),
            "i18n_count": len(t.i18n_fields),
            "has_i18n": t.has_i18n,
        }
        for t in ws.tables
    ]


def _i18n_table(ws: CanonicalWorkspace, table: str) -> object:
    """Find an i18n-capable table by name, raising a friendly error otherwise."""
    _require_table(ws, table)
    t = next((t for t in ws.tables if t.table == table), None)
    if not t.has_i18n:
        raise ValueError(f"表 '{table}' 没有 i18n 字段")
    return t


def canonical_i18n_entries(root: Path, table: str, lang: str) -> list[dict]:
    """Return computed translation entries for a table+lang (source + text + status)."""
    from ct.export.i18n.merger import load_translation
    from ct.export.i18n.state import sync_lang_table

    ws = CanonicalWorkspace.load(root)
    config = ws.config
    target = _i18n_table(ws, table)
    if lang not in config.secondary_langs:
        raise ValueError(f"语言 '{lang}' 不在 secondary_langs 中")
    i18n_dir = config.resolve("i18n_dir")
    source_path = i18n_dir / "source" / f"{table}.json"
    source = (
        json.loads(source_path.read_text(encoding="utf-8"))
        if source_path.exists()
        else {}
    )
    computed = sync_lang_table(source, load_translation(i18n_dir, lang, table))
    entries: list[dict] = []
    for key, entry in computed.items():
        id_part, _, field = key.partition(".")
        entries.append(
            {
                "key": key,
                "id": id_part,
                "field": field,
                "source": str(entry.get("source", "")),
                "text": str(entry.get("text", "")),
                "confirmed": bool(entry.get("confirmed", False)),
                "status": str(entry.get("status", "missing")),
            }
        )
    return entries


def canonical_i18n_save_entry(
    root: Path,
    table: str,
    lang: str,
    key: str,
    text: str,
    confirmed: bool,
) -> dict:
    """Save a single translation entry, recompute its status, and re-dump the lang file."""
    from ct.export.i18n.merger import load_translation, write_lang_file
    from ct.export.i18n.state import compute_status

    ws = CanonicalWorkspace.load(root)
    config = ws.config
    target = _i18n_table(ws, table)
    if lang not in config.secondary_langs:
        raise ValueError(f"语言 '{lang}' 不在 secondary_langs 中")
    i18n_dir = config.resolve("i18n_dir")
    entries = load_translation(i18n_dir, lang, table)
    if key not in entries:
        raise ValueError(f"条目 {key} 不存在，请先同步骨架")
    entries[key]["text"] = str(text)
    entries[key]["confirmed"] = bool(confirmed)
    entries[key]["status"] = compute_status(
        entries[key]["text"], entries[key]["confirmed"], in_source=True
    ).value
    field_order = [f.name for f in target.i18n_fields]
    write_lang_file(entries, i18n_dir / lang / f"{table}.json", field_order)
    return entries[key]
